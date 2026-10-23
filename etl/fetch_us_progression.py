"""U.S. supply & use progression by WASDE release (corn, wheat, soybeans) -> JSA.WASDE.US_PROGRESSION.

Reproduces JSA's "WASDE Progression.xlsx" template: for each marketing year, the estimate at each monthly WASDE
(May of the year it starts through the following April) plus the 'Current' (latest/final) estimate.

  python etl/fetch_us_progression.py --seed-workbook "C:\\...\\WASDE Progression.xlsx"   # history 2011/12-> from the Excel
  python etl/fetch_us_progression.py --since 2025-05                                      # USDA's monthly WASDE CSVs
  python etl/fetch_us_progression.py                                                      # last 3 months (the cron mode)

Units are WASDE's own: million acres, bushels/acre, million bushels, stocks/use in %.
Sources: the workbook (source='excel') and USDA's per-report CSV (source='wasde_csv'; no key, browser User-Agent
required, reports exist from Aug 2021). A WASDE CSV row always wins over the workbook for the same release.
"""
from __future__ import annotations

import argparse
import io
import re
import sys
import time
from datetime import date, datetime
from pathlib import Path

import pandas as pd
import requests

sys.path.insert(0, str(Path(__file__).parent.parent))
sys.path.insert(0, str(Path(__file__).parent))
from fetch_wasde import _sf_connect  # noqa: E402

UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
                    "Chrome/120.0 Safari/537.36"}
CSV_URL = "https://www.usda.gov/sites/default/files/documents/oce-wasde-report-data-{y}-{m:02d}.csv"
COMMODITY_CSV = {"Corn": "Corn", "Wheat": "Wheat", "Soybeans": "Oilseed, Soybean"}

DDL = [
    """CREATE TABLE IF NOT EXISTS JSA.WASDE.US_PROGRESSION (
        commodity VARCHAR, market_year INTEGER, kind VARCHAR, release_date DATE, item VARCHAR, value FLOAT,
        source VARCHAR, loaded_at TIMESTAMP_NTZ DEFAULT CURRENT_TIMESTAMP)""",
    "GRANT SELECT ON TABLE JSA.WASDE.US_PROGRESSION TO ROLE WASDE_ROLE",
    "GRANT SELECT, INSERT, DELETE, TRUNCATE ON TABLE JSA.WASDE.US_PROGRESSION TO ROLE WASDE_ETL_ROLE",
]

# WASDE CSV attribute -> template item (per commodity; values are US, annual, native units)
CSV_ITEMS = {
    "Area Harvested": "Harvested Area", "Yield per Harvested Acre": "Yield", "Beginning Stocks": "Beginning Stocks",
    "Production": "Production", "Imports": "Imports", "Supply, Total": "Total Supply", "Exports": "Exports",
    "Feed and Residual": "Feed", "Food, Seed & Industrial": "FSI", "Domestic, Total": "Domestic Consumption",
    "Use, Total": "Total Demand", "Ending Stocks": "Ending Stocks", "Crushings": "Crush",
}
CSV_UNITS = {"Harvested Area": "Million Acres", "Yield": "Bushels"}          # everything else: Million Bushels

# Excel row label (lower-cased prefix) -> template item
XL_ITEMS = [("harvested area", "Harvested Area"), ("yield", "Yield"), ("beginning stocks", "Beginning Stocks"),
            ("production", "Production"), ("import", "Imports"), ("total supply", "Total Supply"),
            ("crush", "Crush"), ("exports", "Exports"), ("feed", "Feed"), ("fsi", "FSI"), ("food, seed", "FSI"),
            ("domestic consump", "Domestic Consumption"), ("total demand", "Total Demand"),
            ("ending stocks", "Ending Stocks"), ("stocks/use", "Stocks/Use")]


def _xl_item(label) -> str | None:
    t = str(label or "").strip().lower()
    return next((v for k, v in XL_ITEMS if t.startswith(k)), None)


# ── workbook ──────────────────────────────────────────────────────────────────
def parse_workbook(path: Path) -> pd.DataFrame:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = []
    for sheet in wb.sheetnames:
        com = sheet.strip()
        if com not in COMMODITY_CSV:
            continue
        ws = wb[sheet]
        data = [list(r) for r in ws.iter_rows(values_only=True)]
        i = 0
        while i < len(data) - 2:
            hdr, lab = data[i], data[i + 1]
            if isinstance(hdr[1], (datetime, date)) and re.fullmatch(r"\d{4}/\d{2}", str(lab[1] or "").strip()):
                my = int(str(lab[1])[:4])
                dates = {j: hdr[j].date() if isinstance(hdr[j], datetime) else hdr[j]
                         for j in range(1, 13) if isinstance(hdr[j], (datetime, date))}
                k = i + 2
                while k < len(data):
                    name = _xl_item(data[k][0])
                    if name:
                        for j, d in dates.items():
                            v = data[k][j]
                            if isinstance(v, (int, float)):
                                rows.append((com, my, "release", d, name, float(v) * (100 if name == "Stocks/Use" else 1), "excel"))
                        v = data[k][13] if len(data[k]) > 13 else None
                        if isinstance(v, (int, float)):
                            rows.append((com, my, "current", None, name, float(v) * (100 if name == "Stocks/Use" else 1), "excel"))
                    if str(data[k][0] or "").strip().lower().startswith("stocks/use"):
                        break
                    k += 1
                i = k
            i += 1
    return pd.DataFrame(rows, columns=["commodity", "market_year", "kind", "release_date", "item", "value", "source"])


# ── USDA WASDE CSV ────────────────────────────────────────────────────────────
def fetch_report(y: int, m: int) -> pd.DataFrame | None:
    r = requests.get(CSV_URL.format(y=y, m=m), headers=UA, timeout=90)
    if r.status_code == 404:
        return None
    r.raise_for_status()
    d = pd.read_csv(io.BytesIO(r.content), encoding="latin-1")
    rel = pd.to_datetime(d["ReleaseDate"].dropna().iloc[0]).date()
    out = []
    for com, csv_name in COMMODITY_CSV.items():
        x = d[(d.Commodity == csv_name) & (d.Region == "United States")
              & (d.AnnualQuarterFlag.fillna("Annual") == "Annual")]
        for my_label, g in x.groupby("MarketYear"):
            if not re.fullmatch(r"\d{4}/\d{2}", str(my_label)):
                continue
            vals = {}
            for _, r in g.iterrows():
                item = CSV_ITEMS.get(r.Attribute)
                want = CSV_UNITS.get(item, "Million Bushels") if item else None
                if item and r.Unit == want and pd.notna(r.Value):
                    vals[item] = float(r.Value)
            if com == "Wheat":                                       # template FSI = food + seed
                f = g[(g.Attribute == "Food") & (g.Unit == "Million Bushels")].Value
                s = g[(g.Attribute == "Seed") & (g.Unit == "Million Bushels")].Value
                if len(f) and len(s):
                    vals["FSI"] = float(f.iloc[0]) + float(s.iloc[0])
            if com == "Soybeans" and "Total Demand" in vals and "Exports" in vals:   # crush + seed + residual
                vals["Domestic Consumption"] = vals["Total Demand"] - vals["Exports"]
            if "Ending Stocks" in vals and vals.get("Total Demand"):
                vals["Stocks/Use"] = vals["Ending Stocks"] / vals["Total Demand"] * 100
            for item, v in vals.items():
                out.append((com, int(str(my_label)[:4]), "release", rel, item, v, "wasde_csv"))
    return pd.DataFrame(out, columns=["commodity", "market_year", "kind", "release_date", "item", "value", "source"])


# ── ESMIS XML (fallback when USDA's flat CSV for a month is missing) ──────────
ESMIS = "https://esmis.nal.usda.gov/api/v1/release/findByPubId/1659"      # WASDE
XML_TABLES = {                                    # sub-report title -> (commodity, attribute tag of ITS matrix)
    "U.S. Wheat Supply and Use": ("Wheat", "attribute1"),
    "U.S. Feed Grain and Corn Supply and Use": ("Corn", "attribute2"),     # matrix 2 = corn (matrix 1 = all feed grains)
    "U.S. Soybeans and Products Supply and Use": ("Soybeans", "attribute4"),
}


def esmis_xml_urls(since: date) -> dict[tuple[int, int], tuple[date, str]]:
    out, page = {}, 0
    while True:
        r = requests.get(ESMIS, params={"page": page}, timeout=60).json()
        for it in r["results"]:
            d = datetime.fromisoformat(it["release_datetime"].replace("+0000", "+00:00")).date()
            if d < since:
                return out
            xml = next((f for f in it["files"] if f.lower().endswith(".xml")), None)
            if xml:
                out.setdefault((d.year, d.month), (d, xml))
        page += 1
        if page >= r["pager"]["total_pages"]:
            return out


def parse_xml_report(xml_text: str, rel: date) -> pd.DataFrame:
    import xml.etree.ElementTree as ET
    root = ET.fromstring(xml_text.lstrip("﻿"))
    mon = rel.strftime("%b")
    out = []
    for rep in root.iter("Report"):
        title = re.sub(r"\s+", " ", (rep.get("sub_report_title") or "")).strip()
        key = next((k for k in XML_TABLES if title.startswith(k)), None)
        if not key:
            continue
        com, tag = XML_TABLES[key]
        vals: dict[int, dict[str, float]] = {}
        for a in rep.iter(tag):
            name = re.sub(r"\s+\d/$", "", re.sub(r"\s+", " ", a.get(tag, "")).strip()).strip()
            item = CSV_ITEMS.get(name) or ({"Food": "Food", "Seed": "Seed"}.get(name))
            if not item:
                continue
            for yg in a.iter():
                if not yg.tag.endswith("_year_group"):
                    continue
                m = re.match(r"(\d{4})/\d{2}", next((v for k, v in yg.attrib.items() if k.startswith("market_year")), ""))
                if not m:
                    continue
                for mg in yg.iter():
                    if not mg.tag.endswith("_month_group"):
                        continue
                    fm = next((v for k, v in mg.attrib.items() if k.startswith("forecast_month")), "")
                    if fm and fm[:3] != mon:                       # some tables spell the month out
                        continue
                    cell = next((c for c in mg.iter("Cell")), None)
                    raw = next((v for k, v in (cell.attrib.items() if cell is not None else []) if k.startswith("cell_value")), None)
                    num = re.match(r"\s*(-?[\d,]*\.?\d+)", str(raw))             # tolerate footnote marks ("87.4 *")
                    if not num:
                        continue
                    v = float(num.group(1).replace(",", ""))
                    vals.setdefault(int(m.group(1)), {})[item] = v
        for my, v in vals.items():
            if com == "Wheat" and "Food" in v and "Seed" in v:
                v["FSI"] = v["Food"] + v["Seed"]
            if com == "Soybeans" and "Total Demand" in v and "Exports" in v:
                v["Domestic Consumption"] = v["Total Demand"] - v["Exports"]
            if v.get("Ending Stocks") is not None and v.get("Total Demand"):
                v["Stocks/Use"] = v["Ending Stocks"] / v["Total Demand"] * 100
            for item, x in v.items():
                if item in ("Food", "Seed"):
                    continue
                out.append((com, my, "release", rel, item, x, "wasde_xml"))
    return pd.DataFrame(out, columns=["commodity", "market_year", "kind", "release_date", "item", "value", "source"])


def months(since: tuple[int, int], until: tuple[int, int]):
    y, m = since
    while (y, m) <= until:
        yield y, m
        m += 1
        if m > 12:
            y, m = y + 1, 1


def _load(conn, df: pd.DataFrame, delete_sql: str, params: tuple):
    cur = conn.cursor()
    try:
        cur.execute(delete_sql, params)
        cols = ["commodity", "market_year", "kind", "release_date", "item", "value", "source"]
        rows = [tuple(None if pd.isna(v) else (v.item() if hasattr(v, "item") else v) for v in r)
                for r in df[cols].itertuples(index=False, name=None)]
        cur.executemany("INSERT INTO JSA.WASDE.US_PROGRESSION (commodity, market_year, kind, release_date, item, "
                        "value, source) VALUES (%s,%s,%s,%s,%s,%s,%s)", rows)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()


def main():
    from dotenv import load_dotenv
    load_dotenv(Path(__file__).parent.parent / ".env")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--seed-workbook", help="load the history from JSA's WASDE Progression.xlsx")
    ap.add_argument("--since", help="YYYY-MM: fetch every WASDE CSV from this month (default: 3 months back)")
    ap.add_argument("--setup", action="store_true", help="create the table + grants and exit")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    conn = _sf_connect()
    if args.setup or args.seed_workbook:       # table creation needs the owner role; the Droplet's ETL role can't (and needn't)
        cur = conn.cursor()
        for stmt in DDL:
            cur.execute(stmt)
        conn.commit()
        cur.close()
    if args.setup:
        print("Setup done.")
        return

    if args.seed_workbook:
        df = parse_workbook(Path(args.seed_workbook))
        print(f"Workbook: {len(df)} values, {df.commodity.nunique()} commodities, "
              f"marketing years {df.market_year.min()}-{df.market_year.max()}")
        if not args.dry_run:
            _load(conn, df, "DELETE FROM JSA.WASDE.US_PROGRESSION WHERE source = %s", ("excel",))
            print("Workbook loaded.")

    today = date.today()
    if args.since:
        y, m = (int(x) for x in args.since.split("-"))
    else:
        y, m = (today.year, today.month - 2) if today.month > 2 else (today.year - 1, today.month + 10)
    xml_index = None
    for yy, mm in months((y, m), (today.year, today.month)):
        df = fetch_report(yy, mm)
        time.sleep(1)
        if df is None or df.empty:                                  # flat CSV missing: try USDA's archive XML
            if xml_index is None:
                xml_index = esmis_xml_urls(date(y, m, 1))
            hit = xml_index.get((yy, mm))
            if hit:
                df = parse_xml_report(requests.get(hit[1], timeout=120).content.decode("utf-8", "ignore"), hit[0])
                print(f"  {yy}-{mm:02d}: (from archive XML)")
        if df is None or df.empty:
            print(f"  {yy}-{mm:02d}: no report")
            continue
        rel = df.release_date.iloc[0]
        print(f"  {yy}-{mm:02d}: released {rel}, {len(df)} values")
        if not args.dry_run:
            _load(conn, df, "DELETE FROM JSA.WASDE.US_PROGRESSION WHERE source IN ('wasde_csv', 'wasde_xml') "
                  "AND release_date = %s", (rel,))
    conn.close()
    print("Done.")


if __name__ == "__main__":
    main()
