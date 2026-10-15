"""Load JSA's hand-kept CONAB workbooks (soybeans; corn 1st / 2nd / 3rd crop) as the source for the Brazil tab.

Run by hand from a desktop that can see the SharePoint folder (the Droplet can't):
    python etl/load_conab_workbooks.py [--dir <folder>] [--dry-run]

Rows go to JSA.WASDE.CONAB_USER_SERIES, which db.conab_series() prefers over the public CONAB_SERIES for any
(crop, metric, region, year) it holds, so the monthly CONAB reload can never overwrite them. Corn (total) is the
sum of the three crops; wheat and any season the workbooks don't have (e.g. 2026/27) still come from CONAB's files.
"""
from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).parent.parent))
sys.path.insert(0, str(Path(__file__).parent))
from fetch_south_america import CONAB_REGIONS, _replace  # noqa: E402
from fetch_wasde import _sf_connect  # noqa: E402

DEFAULT_DIR = (r"C:\Users\KoltenPostin\John Stewart and Associates\JSA - Documents\Research Analyst\Misc\CONAB")
FILES = {
    "BRAZIL SOYBEAN CROP-CONAB.xlsx": {"Soybeans": 0},
    "BRAZIL CORN CROPS-CONAB.xlsx": {"Corn (1st crop)": 0, "Corn (2nd crop)": 1, "Corn (3rd crop)": 2},
}
SHEETS = {"area_kha": "plantings m ha's", "yield_kgha": "yield-mt per ha", "production_kt": "production"}

# English region / state names in the workbooks -> CONAB's codes used by CONAB_SERIES
NAMES = {
    "north": "NORTE", "roraima": "RR", "rondonia": "RO", "acre": "AC", "amazonas": "AM", "amapa": "AP",
    "para": "PA", "tocantins": "TO", "north east": "NORDESTE", "maranhao": "MA", "piaui": "PI", "ceara": "CE",
    "rio grande do norte": "RN", "paraiba": "PB", "pernambuco": "PE", "alagoas": "AL", "sergipe": "SE",
    "bahia": "BA", "center west": "CENTRO-OESTE", "mato grosso": "MT", "mato grosso do sul": "MS",
    "goias": "GO", "federal district": "DF", "south east": "SUDESTE", "minas gerais": "MG",
    "espinto santo": "ES", "espirito santo": "ES", "rio de janeiro": "RJ", "sao paulo": "SP", "south": "SUL",
    "parana": "PR", "santa catarina": "SC", "rio grande do sul": "RS", "north/ north east": "NORTE/NORDESTE",
    "north/north east": "NORTE/NORDESTE", "center south": "CENTRO-SUL", "brazil": "BRASIL", "brasil": "BRASIL",
}
CODES = set(CONAB_REGIONS) | set(NAMES.values())     # some blocks already use CONAB's own codes (SP, NORTE, ...)
DDL = [
    """CREATE TABLE IF NOT EXISTS JSA.WASDE.CONAB_USER_SERIES (
        crop VARCHAR, metric VARCHAR, region VARCHAR, level VARCHAR,
        start_year INTEGER, label VARCHAR, is_forecast BOOLEAN, value FLOAT, source VARCHAR,
        loaded_at TIMESTAMP_NTZ DEFAULT CURRENT_TIMESTAMP)""",
    "GRANT SELECT ON TABLE JSA.WASDE.CONAB_USER_SERIES TO ROLE WASDE_ROLE",
    "GRANT SELECT, INSERT, DELETE, TRUNCATE ON TABLE JSA.WASDE.CONAB_USER_SERIES TO ROLE WASDE_ETL_ROLE",
]


def _norm(s) -> str:
    t = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", t).strip().lower()


def _blocks(rows):
    """Split a sheet into REGION/STATE blocks (corn stacks 1st, 2nd, 3rd crop); each runs through BRAZIL."""
    out, cur = [], None
    for r in rows:
        lab = _norm(r[0]) if r[0] is not None else ""
        if lab == "region/state":
            cur = []
            out.append(cur)
        elif cur is not None and lab:
            cur.append(r)
            if lab in ("brazil", "brasil"):
                cur = None
    return out


def parse_workbook(path: Path, crops: dict) -> pd.DataFrame:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = {_norm(n): n for n in wb.sheetnames}
    recs = []
    for metric, sname in SHEETS.items():
        rows = list(wb[sheets[sname]].iter_rows(values_only=True))
        header = next(r for r in rows if r[0] is not None and _norm(r[0]) == "region/state")   # first block's header
        cols = {}
        for j in range(1, len(header)):
            m = re.match(r"\s*(\d{4})\s*/\s*(\d{2})", str(header[j])) if header[j] is not None else None
            if m:
                cols[j] = (int(m.group(1)), f"{m.group(1)}/{m.group(2)}")
        blocks = _blocks(rows)
        for crop, bi in crops.items():
            for r in blocks[bi]:
                name = _norm(r[0])
                code = NAMES.get(name) or (name.upper() if name.upper() in CODES else None)
                if code is None:
                    raise KeyError(f"{path.name} / {sname}: unmapped region label {r[0]!r}")
                level = "total" if code == "BRASIL" else "region" if code in CONAB_REGIONS else "state"
                for j, (sy, label) in cols.items():
                    v = r[j] if j < len(r) else None
                    if isinstance(v, (int, float)) and not pd.isna(v) and v != 0:
                        recs.append((crop, metric, code, level, sy, label, False, float(v), path.name))
    return pd.DataFrame(recs, columns=["crop", "metric", "region", "level", "start_year", "label",
                                       "is_forecast", "value", "source"])


def add_corn_total(df: pd.DataFrame) -> pd.DataFrame:
    """Corn (total) = sum of the three crops; yield recomputed from production / area."""
    c = df[df.crop.str.startswith("Corn (")]
    key = ["region", "level", "start_year", "label", "is_forecast", "source"]
    s = (c[c.metric.isin(["area_kha", "production_kt"])]
         .groupby(["metric"] + key, as_index=False).value.sum())
    piv = s.pivot_table(index=key, columns="metric", values="value").reset_index()
    piv["yield_kgha"] = piv.production_kt / piv.area_kha * 1000
    rows = piv.melt(id_vars=key, value_vars=["area_kha", "production_kt", "yield_kgha"],
                    var_name="metric", value_name="value").dropna(subset=["value"])
    rows["crop"] = "Corn"
    return pd.concat([df, rows[df.columns]], ignore_index=True)


def main():
    from dotenv import load_dotenv
    load_dotenv(Path(__file__).parent.parent / ".env")
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", default=DEFAULT_DIR)
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    frames = []
    for fname, crops in FILES.items():
        df = parse_workbook(Path(args.dir) / fname, crops)
        print(f"  {fname}: {len(df)} rows, crops {sorted(df.crop.unique())}, "
              f"seasons {df.start_year.min()}-{df.start_year.max()}")
        frames.append(df)
    out = add_corn_total(pd.concat(frames, ignore_index=True))
    print(f"  total {len(out)} rows (incl. Corn total)")
    if args.dry_run:
        print("Dry run: nothing written.")
        return
    conn = _sf_connect()
    cur = conn.cursor()
    for stmt in DDL:
        cur.execute(stmt)
    conn.commit()
    _replace(conn, "CONAB_USER_SERIES", out,
             ["crop", "metric", "region", "level", "start_year", "label", "is_forecast", "value", "source"])
    conn.close()
    print("Done.")


if __name__ == "__main__":
    main()
